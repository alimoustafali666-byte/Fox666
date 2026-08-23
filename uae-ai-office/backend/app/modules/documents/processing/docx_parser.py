"""DOCX extraction: paragraph order, headings, and tables. No pixel-
perfect formatting reconstruction -- text content and coarse structure
only.
"""

from typing import BinaryIO

import docx

from app.core.config import settings
from app.modules.documents.processing.base import DocumentParser, ParsedDocument, ParsedElement
from app.modules.documents.processing.errors import InsufficientTextError, ParserFailureError
from app.modules.documents.processing.zip_safety import assert_zip_is_safe_to_decompress


class DocxParser(DocumentParser):
    def supports(self, file_type: str) -> bool:
        return (
            file_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    def parse(self, fileobj: BinaryIO) -> ParsedDocument:
        assert_zip_is_safe_to_decompress(
            fileobj,
            max_entry_bytes=settings.processing_max_zip_entry_uncompressed_bytes,
            max_total_bytes=settings.processing_max_zip_total_uncompressed_bytes,
        )
        fileobj.seek(0)

        try:
            document = docx.Document(fileobj)

            elements: list[ParsedElement] = []
            current_section: str | None = None

            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = (paragraph.style.name if paragraph.style else "") or ""
                is_heading = style_name.lower().startswith("heading") or style_name.lower() == "title"
                if is_heading:
                    current_section = text
                    elements.append(
                        ParsedElement(
                            text=text,
                            element_type="heading",
                            section_name=current_section,
                            source_location={"section": current_section},
                        )
                    )
                else:
                    elements.append(
                        ParsedElement(
                            text=text,
                            element_type="paragraph",
                            section_name=current_section,
                            source_location={"section": current_section} if current_section else {},
                        )
                    )

            for table_index, table in enumerate(document.tables):
                for row_index, row in enumerate(table.rows):
                    cells = [cell.text.strip() for cell in row.cells]
                    if not any(cells):
                        continue
                    elements.append(
                        ParsedElement(
                            text=" | ".join(cells),
                            element_type="table_row",
                            section_name=current_section,
                            source_location={"table_index": table_index, "row_index": row_index},
                        )
                    )
        except Exception as exc:
            raise ParserFailureError("DOCX could not be read as a valid file.") from exc

        if not elements:
            raise InsufficientTextError("DOCX has no extractable text content.")

        return ParsedDocument(elements=elements)

