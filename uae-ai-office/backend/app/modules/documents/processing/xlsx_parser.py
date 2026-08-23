"""XLSX extraction: a deterministic, structural textual representation
-- sheet name, row relationships, and cell values, using the first
non-empty row of each sheet as a header to label subsequent rows'
values (no column names are hard-coded; this works for any workbook
shape). Formula cells are read via their last cached calculated value
(openpyxl's data_only=True) -- formulas are never evaluated by this
code, and macros are never executed (XLSM stays unsupported entirely,
rejected already at Step 8's upload validation).
"""

from datetime import date, datetime
from typing import BinaryIO

import openpyxl
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.modules.documents.processing.base import DocumentParser, ParsedDocument, ParsedElement
from app.modules.documents.processing.errors import (
    InsufficientTextError,
    ParserFailureError,
    PathologicalDocumentError,
)
from app.modules.documents.processing.zip_safety import assert_zip_is_safe_to_decompress


def _stringify_cell_value(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _is_blank(value: object) -> bool:
    return value is None or str(value).strip() == ""


class XlsxParser(DocumentParser):
    def supports(self, file_type: str) -> bool:
        return file_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def parse(self, fileobj: BinaryIO) -> ParsedDocument:
        assert_zip_is_safe_to_decompress(
            fileobj,
            max_entry_bytes=settings.processing_max_zip_entry_uncompressed_bytes,
            max_total_bytes=settings.processing_max_zip_total_uncompressed_bytes,
        )
        fileobj.seek(0)

        try:
            workbook = openpyxl.load_workbook(fileobj, data_only=True, read_only=True)
        except Exception as exc:
            raise ParserFailureError("XLSX could not be read as a valid file.") from exc

        try:
            sheet_names = workbook.sheetnames
            if len(sheet_names) > settings.processing_max_xlsx_sheets:
                raise PathologicalDocumentError(
                    f"Workbook has more than {settings.processing_max_xlsx_sheets} sheets."
                )

            elements: list[ParsedElement] = []
            for sheet_name in sheet_names:
                elements.extend(self._parse_sheet(workbook[sheet_name], sheet_name))
        except PathologicalDocumentError:
            raise
        except Exception as exc:
            raise ParserFailureError("XLSX could not be read as a valid file.") from exc
        finally:
            workbook.close()

        if not elements:
            raise InsufficientTextError("Workbook has no extractable cell content.")

        return ParsedDocument(elements=elements)

    def _parse_sheet(self, sheet, sheet_name: str) -> list[ParsedElement]:
        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        if max_row > settings.processing_max_xlsx_rows_per_sheet:
            raise PathologicalDocumentError(
                f"Sheet '{sheet_name}' has more than "
                f"{settings.processing_max_xlsx_rows_per_sheet} rows."
            )
        if max_col > settings.processing_max_xlsx_columns_per_sheet:
            raise PathologicalDocumentError(
                f"Sheet '{sheet_name}' has more than "
                f"{settings.processing_max_xlsx_columns_per_sheet} columns."
            )

        elements: list[ParsedElement] = []
        headers: list[str | None] = []
        header_found = False

        for row_index, row in enumerate(sheet.iter_rows(values_only=False), start=1):
            # Position-derived column letters, not cell.column_letter --
            # a row's trailing cells beyond the sheet's populated area
            # can come back as bare EmptyCell objects in read-only mode,
            # which don't carry that attribute.
            cells = [(get_column_letter(idx), cell.value) for idx, cell in enumerate(row, start=1)]
            if all(_is_blank(value) for _, value in cells):
                continue  # skip fully empty rows

            if not header_found:
                headers = [
                    None if _is_blank(value) else _stringify_cell_value(value).strip()
                    for _, value in cells
                ]
                header_found = True
                continue

            lines = []
            for idx, (col_letter, value) in enumerate(cells):
                if _is_blank(value):
                    continue
                label = headers[idx] if idx < len(headers) and headers[idx] else col_letter
                lines.append(f"{label}: {_stringify_cell_value(value)}")

            if not lines:
                continue

            elements.append(
                ParsedElement(
                    text=f"Row {row_index}:\n" + "\n".join(lines),
                    element_type="table_row",
                    sheet_name=sheet_name,
                    source_location={"sheet_name": sheet_name, "row": row_index},
                )
            )

        return elements

